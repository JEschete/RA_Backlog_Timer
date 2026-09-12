"""Excel and CSV export.

Excel is now an output format rather than the live database, so it is free to
be formatted for reading: frozen header, autofilter, sensible widths, a colour
scale on the efficiency column and a summary sheet.
"""
from __future__ import annotations

import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..config import log
from ..efficiency import add_metrics, summarize
from . import repo

# Display header -> source column.
_EXPORT_COLUMNS = [
    ("Title", "title"),
    ("System", "system"),
    ("Achievements", "achievements"),
    ("Earned", "earned"),
    ("Points", "points"),
    ("RA_ID", "ra_id"),
    ("HLTB_Beat", "hltb_beat"),
    ("HLTB_Complete", "hltb_complete"),
    ("RA_Beat", "ra_beat"),
    ("RA_Master", "ra_master"),
    ("RA_Beat_HC", "ra_beat_hardcore"),        # item #11
    ("RA_Master_HC", "ra_master_hardcore"),    # item #11
    ("RA_Players", "ra_players"),
    ("Points_Per_Hour", "points_per_hour"),
    ("Remaining_Hours", "remaining_hours"),
    ("Match_Quality", "match_quality"),
    ("HLTB_Name", "hltb_name"),
]

_WIDTHS = {
    "Title": 42, "System": 22, "HLTB_Name": 34, "Match_Quality": 14,
}

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def build_export_frame(conn: sqlite3.Connection) -> pd.DataFrame:
    df = add_metrics(repo.load_games(conn))
    out = pd.DataFrame()
    for header, source in _EXPORT_COLUMNS:
        out[header] = df[source] if source in df.columns else None
    return out.sort_values("Points_Per_Hour", ascending=False, na_position="last")


def to_csv(conn: sqlite3.Connection, path: str | Path) -> Path:
    path = Path(path).with_suffix(".csv")
    build_export_frame(conn).to_csv(path, index=False, encoding="utf-8-sig")
    log.info("Exported CSV to %s", path)
    return path


def to_excel(conn: sqlite3.Connection, path: str | Path) -> Path:
    """Write a formatted workbook.

    Raises PermissionError to the caller if the file is open elsewhere -- the
    caller decides what to do, rather than the scan dying mid-run (item #6).
    """
    path = Path(path).with_suffix(".xlsx")
    df = build_export_frame(conn)
    stats = summarize(add_metrics(repo.load_games(conn)))

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Backlog", index=False)
        _summary_frame(stats, conn).to_excel(
            writer, sheet_name="Summary", index=False)
        _format_backlog(writer.book["Backlog"], df)
        _format_summary(writer.book["Summary"])

    log.info("Exported Excel to %s", path)
    return path


def _summary_frame(stats: dict, conn: sqlite3.Connection) -> pd.DataFrame:
    rows = [
        ("Total games", stats["total_games"]),
        ("With HLTB data", stats["with_hltb"]),
        ("With RA mastery data", stats["with_ra_mastery"]),
        ("Without any time data", stats["no_time_data"]),
        ("Total points", stats["total_points"]),
        ("Total hours", stats["total_hours"]),
        ("Total days", round(stats["total_hours"] / 24, 1)),
        ("Average hours per game", stats["avg_hours"]),
        ("", ""),
    ]
    rows += [("System: " + name, count) for name, count in repo.list_systems(conn)]
    return pd.DataFrame(rows, columns=["Metric", "Value"])


def _format_backlog(sheet, df: pd.DataFrame) -> None:
    last_row = sheet.max_row
    last_col = sheet.max_column

    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(last_col)}{last_row}"

    for idx, (header, _) in enumerate(_EXPORT_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = _WIDTHS.get(header, 14)

    if last_row < 2:
        return

    # Green = more points per hour of your life. That is the whole point of
    # the column, so make it readable at a glance.
    if "Points_Per_Hour" in df.columns:
        col = get_column_letter(df.columns.get_loc("Points_Per_Hour") + 1)
        sheet.conditional_formatting.add(
            f"{col}2:{col}{last_row}",
            ColorScaleRule(
                start_type="min", start_color="F8696B",
                mid_type="percentile", mid_value=50, mid_color="FFEB84",
                end_type="max", end_color="63BE7B",
            ),
        )

    for header in ("HLTB_Beat", "HLTB_Complete", "RA_Beat", "RA_Master",
                   "RA_Beat_HC", "RA_Master_HC", "Points_Per_Hour",
                   "Remaining_Hours"):
        if header not in df.columns:
            continue
        col = get_column_letter(df.columns.get_loc(header) + 1)
        for cell in sheet[col][1:]:
            cell.number_format = "0.0"


def _format_summary(sheet) -> None:
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 16
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
